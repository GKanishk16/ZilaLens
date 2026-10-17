"""Create map previews and Power BI/Excel handoff assets."""
from pathlib import Path
import json,re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib import colors
ROOT=Path(__file__).resolve().parents[1]

def norm(s): return ' '.join(re.sub(r'[^a-z0-9 ]',' ',str(s).lower()).split())
def _polygons(geom):
    if not geom: return []
    if geom['type']=='Polygon': return geom['coordinates']
    if geom['type']=='MultiPolygon': return [p for p in geom['coordinates']]
    return []

def make_map(panel, indicator, title, path, state=None):
    geo=json.loads((ROOT/'data'/'raw'/'india_districts.geojson').read_text())
    data=panel[(panel['round']=='NFHS-4')&(panel['indicator']==indicator)].copy()
    if state: data=data[data.state.map(norm)==norm(state)]
    vals={(norm(r.district),norm(r.state)):r.value for r in data.itertuples() if pd.notna(r.value)}
    fig,ax=plt.subplots(figsize=(7,8)); cmap=plt.get_cmap('YlGnBu'); normc=colors.Normalize(vmin=min(vals.values()) if vals else 0,vmax=max(vals.values()) if vals else 1)
    joined=0
    for feat in geo['features']:
        prop=feat['properties']; district=prop.get('district',''); st=prop.get('st_nm','')
        value=vals.get((norm(district),norm(st)))
        if value is not None: joined+=1
        fc=cmap(normc(value)) if value is not None else '#eeeeea'
        for poly in _polygons(feat['geometry']):
            ring=poly[0]
            ax.fill([p[0] for p in ring],[p[1] for p in ring],facecolor=fc,edgecolor='white',linewidth=.15)
    if state:
        ax.set_xlim(77,85); ax.set_ylim(23,31)
    else: ax.set_xlim(67,98); ax.set_ylim(6,38)
    ax.set_axis_off(); ax.set_title(title,loc='left',fontsize=13,weight='bold',pad=12)
    sm=plt.cm.ScalarMappable(cmap=cmap,norm=normc); fig.colorbar(sm,ax=ax,shrink=.6,label='Percent (%)')
    fig.text(.02,.018,f'n matched districts: {joined} | Source: NFHS-4 (2015-16), HindustanTimesLabs/nfhs-data; boundaries: HindustanTimesLabs/shapefiles, India district GeoJSON.',fontsize=7,wrap=True)
    fig.savefig(path,dpi=200,bbox_inches='tight'); plt.close(fig)
    return joined

def build_powerbi(panel):
    out=ROOT/'outputs'/'powerbi'; out.mkdir(parents=True,exist_ok=True)
    district_cols=['district_id','state','district','aspirational_flag']
    district_cols += [c for c in ['geography_status','geography_eligible','comparison_eligible'] if c in panel.columns]
    panel[district_cols].groupby('district_id',as_index=False,dropna=False).first()[district_cols].to_csv(out/'dim_district.csv',index=False)
    pd.DataFrame({'round_key':[4,5],'round':['NFHS-4','NFHS-5'],'period':['2015-16','2019-21']}).to_csv(out/'dim_round.csv',index=False)
    fact=panel[['district_id','round','indicator','value','source_file']].copy()
    indicator_keys=panel[['indicator']].drop_duplicates().reset_index(drop=True)
    indicator_keys['indicator_key']=indicator_keys.index+1
    indicator_keys.to_csv(out/'dim_indicator.csv',index=False)
    fact=fact.merge(indicator_keys,on='indicator',how='left')
    index_path=out.parent/'tables'/'empowerment_index.csv'
    if index_path.exists():
        try: idx=pd.read_csv(index_path)
        except pd.errors.EmptyDataError: idx=pd.DataFrame()
        if {'district_id','round','index_score'}.issubset(idx.columns): fact=fact.merge(idx[['district_id','round','index_score']],on=['district_id','round'],how='left')
        else: fact['index_score']=np.nan
    else: fact['index_score']=np.nan
    fact.to_csv(out/'fact_indicator_values.csv',index=False)
    cluster_path=out.parent/'tables'/'district_clusters.csv'
    if cluster_path.exists():
        try: pd.read_csv(cluster_path).to_csv(out/'dim_cluster.csv',index=False)
        except pd.errors.EmptyDataError: pd.DataFrame(columns=['district_id','cluster_id','cluster_name']).to_csv(out/'dim_cluster.csv',index=False)
    else: pd.DataFrame(columns=['district_id','cluster_id','cluster_name']).to_csv(out/'dim_cluster.csv',index=False)
    dax=['PP Change = CALCULATE([NFHS-5 Value] - [NFHS-4 Value], dim_district[comparison_eligible] = TRUE())','NFHS-4 Value = CALCULATE(AVERAGE(fact_indicator_values[value]), dim_round[round] = "NFHS-4")','NFHS-5 Value = CALCULATE(AVERAGE(fact_indicator_values[value]), dim_round[round] = "NFHS-5")','District Count = CALCULATE(DISTINCTCOUNT(dim_district[district_id]), dim_district[comparison_eligible] = TRUE())','Districts Improved = CALCULATE(DISTINCTCOUNT(dim_district[district_id]), FILTER(VALUES(dim_district[district_id]), [PP Change] > 0))','Districts Improved % = DIVIDE([Districts Improved], [District Count])','Index Score = AVERAGE(fact_indicator_values[index_score])','Aspirational Mean Change = CALCULATE([PP Change], dim_district[aspirational_flag] = TRUE())','Non-Aspirational Mean Change = CALCULATE([PP Change], dim_district[aspirational_flag] = FALSE())','Aspirational Gap = [Aspirational Mean Change] - [Non-Aspirational Mean Change]']
    (out/'powerbi_build_guide.md').write_text('# Power BI build guide\n\n## Relationships\n\n- `dim_district[district_id]` 1:* `fact_indicator_values[district_id]`\n- `dim_indicator[indicator_key]` 1:* `fact_indicator_values[indicator_key]`\n- `dim_round[round]` 1:* fact\n- `dim_cluster[district_id]` 1:* district (555 screened matched districts)\n\n## Measures\n\n'+''.join(f'{i+1}. `{m}`\n\n' for i,m in enumerate(dax))+'## Pages\n\n1. National Overview: map, indicator selector, round values, distribution and improved share.\n2. Aspirational Districts Comparison: descriptive changes and matched baseline subgroup; label as non-causal.\n3. Uttar Pradesh Deep Dive: district ranks, Ghaziabad/Gautam Buddha Nagar callouts, state/national reference lines.\n\nThe district dimension includes `geography_status`, `geography_eligible` and `comparison_eligible`; the last flag means the ID is present in both rounds and not listed in Harvard Appendix A. The PP Change measure filters to those directly comparable districts. `Index Score` remains blank because both factsheet sources lack bank-use/mobile-use indicators.\n')

def build_workbook(panel):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import BarChart,Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl import load_workbook
    out=ROOT/'outputs'/'tables'/'ZilaLens_Analysis_Workbook.xlsx'
    data=panel.copy()
    summary=panel[panel['round']=='NFHS-4'].groupby(['state','indicator'],as_index=False).value.mean()
    summary['source_file']='data/raw/nfhs_district-wise.csv'
    changes=pd.read_csv(ROOT/'outputs'/'tables'/'district_changes.csv')
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        data.to_excel(writer,sheet_name='Data',index=False)
        summary.to_excel(writer,sheet_name='NFHS4 Summary',index=False)
        changes.to_excel(writer,sheet_name='Change Summary',index=False)
        if len(summary):
            selected=summary[summary.indicator==summary.indicator.iloc[0]][['state','value','source_file']].rename(columns={'value':'mean_percent'})
            selected.to_excel(writer,sheet_name='Chart Data',index=False)
    wb=load_workbook(out);wb.properties.creator='Kanishk Gupta';wb.properties.title='ZilaLens Analysis Workbook'
    for name in wb.sheetnames:
        ws=wb[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True,vertical='center')
        widths={
            'Data':[24,30,12,68,12,48,46,18],
            'NFHS4 Summary':[24,68,14,48],
            'Change Summary':[46,24,30,68,12,12,14,14,20],
            'Chart Data':[24,16,48],
        }.get(name,[])
        for i,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
        if ws.max_row>1:
            for row in ws.iter_rows(min_row=2):
                for c in row:c.alignment=Alignment(vertical='top',wrap_text=False)
                ws.row_dimensions[row[0].row].height=19
    ws=wb['NFHS4 Summary']
    if ws.max_row>1:ws.conditional_formatting.add(f'C2:C{ws.max_row}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
    if wb['Change Summary'].max_row==1: wb['Change Summary']['A2']='No two-round changes calculated: official NFHS-5 workbook is missing.'
    if 'Chart Data' in wb.sheetnames and wb['Chart Data'].max_row>1:
        cw=wb['Chart Data'];chart=BarChart();chart.title=f'NFHS-4 district mean by state: {summary.indicator.iloc[0]} (n={len(selected)} states)';chart.y_axis.title='Mean percent';chart.x_axis.title='State';chart.add_data(Reference(cw,min_col=2,min_row=1,max_row=cw.max_row),titles_from_data=True);chart.set_categories(Reference(cw,min_col=1,min_row=2,max_row=cw.max_row));chart.height=9;chart.width=18
        if chart.series: chart.series[0].graphicalProperties.solidFill='168C8C';chart.series[0].graphicalProperties.line.solidFill='168C8C'
        cw.add_chart(chart,'D2')
    wb.save(out)

if __name__=='__main__':
    p=pd.read_csv(ROOT/'data'/'processed'/'tidy_panel.csv'); build_powerbi(p); build_workbook(p)
    for ind,slug in [('Women who are literate (%)','women_literate'),('Households using improved sanitation facility','improved_sanitation')]:
        if ind in set(p.indicator): make_map(p,ind,'NFHS-4 baseline: '+ind,ROOT/'outputs'/'maps'/f'{slug}_nfhs4.png')
