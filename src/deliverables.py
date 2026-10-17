"""Create client-facing interim deliverables with the current data limitations explicit."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches,Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
ROOT=Path(__file__).resolve().parents[1]
TITLE='ZilaLens'
SUB='Women’s Financial Inclusion and Empowerment Across Indian Districts, NFHS-4 (2015-16) to NFHS-5 (2019-21)'
AUTHOR='Kanishk Gupta'
GITHUB='https://github.com/GKanishk16'
NAVY='17365D'; TEAL='168C8C'; GOLD='C47F17'; MUTED='5B6573'

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def page_field(paragraph):
    run=paragraph.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def style_doc(doc,kind):
    sec=doc.sections[0]; sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor.from_string('263238');normal.paragraph_format.space_after=Pt(6)
    for name,size,color in [('Title',30,NAVY),('Heading 1',19,NAVY),('Heading 2',13,TEAL),('Heading 3',11,NAVY)]:
        st=doc.styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(color)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('ZilaLens | Kanishk Gupta     ');page_field(footer)
    props=doc.core_properties;props.author=AUTHOR;props.title=f'{TITLE} — {kind}';props.subject=SUB;props.keywords='NFHS-4, NFHS-5, district, financial inclusion'

def add_title(doc,tag=None):
    p=doc.add_paragraph();p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(TITLE)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(SUB);r.bold=True;r.font.size=Pt(15);r.font.color.rgb=RGBColor.from_string(TEAL)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(AUTHOR);r.bold=True;r.font.size=Pt(12)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(GITHUB).font.color.rgb=RGBColor.from_string(TEAL)
    if tag:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(tag.upper());r.bold=True;r.font.color.rgb=RGBColor.from_string(GOLD)

def bullet(doc,text): doc.add_paragraph(text,style='List Bullet')
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Light Shading Accent 1'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i];c.text=h;shade(c,NAVY)
        for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):cells[i].text=str(val);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(8.5)
    return t

def report():
    d=Document();style_doc(d,'Research Report');add_title(d,'Interim research report | data-dependent findings pending');d.add_paragraph('Interim edition • 18 October 2026',style='Subtitle');d.add_paragraph('This report records the available evidence and analysis status. NFHS-5 district data are required before the two-round findings can be completed.',style='Intense Quote');d.add_page_break()
    d.add_heading('Contents',0)
    for x in ['Executive summary','Background and research questions','Data and methods','District harmonisation and quality','RQ1 Change in financial inclusion','RQ2 Fast improvers and district typologies','RQ3 Aspirational Districts comparison','RQ4 Uttar Pradesh deep dive','RQ5 Policy evidence','Discussion and recommendations','Limitations','References','Annexes']:bullet(d,x)
    d.add_page_break()
    pages=[
    ('Executive summary',['ZilaLens compares women’s financial inclusion and related empowerment indicators across Indian districts between NFHS-4 (2015-16) and NFHS-5 (2019-21). “Zila” means district in Hindi; the project views women’s financial inclusion through a district lens.','The NFHS-4 source file was downloaded and parsed. Its repository documentation describes 93 district-level indicators; the CSV contains 59,241 rows. (Source: HindustanTimesLabs/nfhs-data, README.md and nfhs_district-wise.csv.)','The official data.gov.in NFHS-5 workbook resource returned HTTP 403 in this environment. No round-to-round change, index, ranking, cluster, subgroup comparison, or UP change finding is reported. The exact workbook and destination are recorded in README.md and docs/DATA_SOURCES.md.','A NITI Aayog list of 112 Aspirational Districts and two public policy documents were acquired. Exact-name matching identifies 96 listed districts in the NFHS-4 data; 16 names remain unmatched and require geography review. (Source files: aspirational_112.pdf and nfhs_district-wise.csv.)','Interim outputs include a reproducible data pipeline, NFHS-4 maps, a bilingual field-validation form, a theory of change, a draft logframe, and a preliminary policy review. These are not substitutes for the pending comparative findings.']),
    ('Background and research questions',['Women’s bank account ownership does not by itself establish active use, control over transactions, or the ability to use digital services safely. The proposed comparison focuses on indicators that are present in both NFHS rounds and retains their source definitions.','The research questions cover district change, fast improvers and lagging districts, descriptive differences between Aspirational and comparable non-aspirational districts, Uttar Pradesh district patterns, and policy drivers and barriers.','NFHS district estimates are survey estimates. District boundaries and names changed between rounds, so comparisons require an explicit crosswalk. This interim report distinguishes direct source observations, computed project status, and unanswered questions.']),
    ('Data and methods',['NFHS-4 district factsheet data were downloaded from HindustanTimesLabs/nfhs-data. The source file contains district, state, census codes, total/rural/urban estimates, indicator number, indicator name, and category. (Source: data/raw/nfhs_district-wise.csv.)','The NFHS-5 catalog resource is an Excel workbook titled “India Districts Factsheets of National Family Health Survey (NFHS) - 5, 2019-2021.” Its direct URL returned HTTP 403 on 18 October 2026. The planned local path is data/raw/NFHS_5_India_Districts_Factsheet_Data.xls.','The pipeline reshapes round data to district, state, round, indicator, and value. It creates normalized name identifiers and records fuzzy name suggestions without applying them. Manual overrides must be documented before any parent/child crosswalk is treated as final.','When both rounds are available, changes will be reported in percentage points and as relative change. Index construction will use common indicators, min-max scaling, equal weights, and sensitivity checks. Any Aspirational District comparison will be described as non-causal.']),
    ('District harmonisation and data quality',['The NFHS-4 file has 59,241 district-indicator rows and 93 district-level indicators according to its source README. The parsed panel currently contains NFHS-4 only. (Source: HindustanTimesLabs/nfhs-data, README.md and nfhs_district-wise.csv.)','The crosswalk currently records normalized-name identifiers and presence by round. It does not yet resolve district splits or mergers. The 2022 NITI 112-district list is an independent source list; exact-name matching finds 96 of its district names in the NFHS-4 source. (Sources: data/processed/district_crosswalk.csv; data/processed/aspirational_districts_112.csv.)','The current GeoJSON contains 645 features. Baseline maps join 563 NFHS-4 district estimates by normalized district and state name. (Source: data/raw/india_districts.geojson; map join counts in figure notes.) This geography is illustrative and not confirmed as boundary-equivalent to the survey rounds.']),
    ('RQ1 Change in financial inclusion',['No two-round changes are available yet. The NFHS-5 workbook is required to identify indicators shared by both rounds and calculate changes using source estimates.','The NFHS-4 compilation includes women’s literacy, 10 or more years of schooling, marriage before age 18, clean cooking fuel, and improved sanitation. The requested bank account and mobile phone measures do not appear under the expected wording in the current compilation. Exact indicator comparability will be recorded after NFHS-5 parsing.','No national, state, or district change statistic is presented here. This prevents a one-round level from being misread as a trend.']),
    ('RQ2 Fast improvers and typologies',['Fastest improvers, lagging districts, and clusters are pending. They require a reviewed district crosswalk and paired values for common indicators.','The analysis code is prepared to calculate district-level absolute and percentage changes, rank districts, and compare baseline level with change. Cluster count selection will use silhouette score when there are sufficient complete district cases.','No cluster names or rank labels are assigned without computed inputs.']),
    ('RQ3 Aspirational Districts comparison',['The requested design is a descriptive difference-in-differences comparison, not a causal evaluation. It will report average changes for Aspirational Districts and other districts, a baseline-matched comparison, and bootstrap intervals.','The NITI Aayog PDF provides the requested 112-district roster. It was published in 2022 and is used here as the project’s specified cohort. The exact-match roster coverage is 96 of 112 NFHS-4 district names; unmatched names require boundary and naming review. (Sources: data/raw/docs/aspirational_112.pdf; data/processed/aspirational_districts_112.csv.)','No comparison estimate or confidence interval is reported without NFHS-5 values and an approved cohort crosswalk. Selection into the programme limits comparability even after baseline matching.']),
    ('RQ4 Uttar Pradesh deep dive',['The UP change ranking is pending NFHS-5. Ghaziabad and Gautam Buddha Nagar will be highlighted against state and national averages only after comparable round estimates are computed.','The project includes interim NFHS-4 district maps for women’s literacy and improved sanitation. These maps are baseline context and do not answer the change question. Figure 1. NFHS-4 women literacy estimates by district. Source: NFHS-4 factsheet compilation; 563 geometry matches; see outputs/maps/women_literate_nfhs4.png.','The final deep dive will show all UP districts, clearly distinguish boundary-matched from unmatched geographies, and cite each statistic to the relevant source table.']),
    ('RQ5 Public policy evidence',['The review covers the PMJDY Mission Document/Brochure and the NITI Aayog/UNDP Aspirational Districts Programme appraisal. Seven paraphrased passages with physical page references are recorded in the coding workbook.','The PMJDY brochure describes expansion of branches and Bank Mitra service points (physical p. 18), transaction interoperability and access (p. 20), and unreliable telecom networks in tribal and hilly areas (p. 30). These are policy design statements, not district outcome estimates.','The ADP appraisal discusses district technical and administrative capacity (physical p. 15), the programme’s convergence, collaboration and competition model (p. 20), and its use of pre-pandemic observations in its own analysis (p. 28). (Sources: data/raw/docs/pmjdy_brochure.pdf; data/raw/docs/adp_appraisal.pdf.)']),
    ('Discussion and recommendations',['The available sources point to a useful distinction between access infrastructure and actual use. The proposed field instrument asks about personal account use, transaction frequency, decision control, phone access, service reliability, and support needs.','For programme planning, first confirm the district-level pattern; then validate whether low reported use reflects account access, service reliability, skills, safety, or household decision constraints. Avoid selecting districts from a name-only comparison.','Build the final intervention around local service mapping, practical assisted account use, digital safety, grievance referral, and routine monitoring. Treat these as design options pending district validation, not proven effects.']),
    ('Limitations',['NFHS-5 access: the official catalog workbook returned HTTP 403. Until it is placed in data/raw, the trend questions remain unanswered.','Aspirational District selection: districts were selected because of pre-existing deprivation and differ from other districts. Descriptive differences and baseline matching cannot remove all selection bias; results will not be interpreted causally.','Boundary changes: NFHS-4 and NFHS-5 districts may have different boundaries. Parent/child mappings can affect estimates and require a documented crosswalk.','Survey sampling error: district estimates are subject to sampling variability. The available factsheet estimates do not provide a complete basis for conventional significance testing of all district changes.','NFHS-5 fieldwork occurred during the COVID-19 period (2019-21), which may affect timing and interpretation. Policy documents also reflect their publication period and do not measure current local service quality.']),
    ('References',['International Institute for Population Sciences (IIPS) and Ministry of Health and Family Welfare (MoHFW). National Family Health Survey-4 district factsheet data, 2015-16. Dataset compilation by HindustanTimesLabs, data/raw/nfhs_district-wise.csv; https://github.com/HindustanTimesLabs/nfhs-data.','NITI Aayog. List of 112 Aspirational Districts. 2022. data/raw/docs/aspirational_112.pdf; https://www.niti.gov.in/sites/default/files/2022-09/List-of-Aspirational-Districts.pdf.','Department of Financial Services, Ministry of Finance. PMJDY Mission Document/Brochure. data/raw/docs/pmjdy_brochure.pdf; https://pmjdy.gov.in/files/E-Documents/PMJDY_BROCHURE_ENG.pdf.','NITI Aayog and United Nations Development Programme. Aspirational Districts Programme: An Appraisal. data/raw/docs/adp_appraisal.pdf; https://www.niti.gov.in/sites/default/files/2023-03/Aspirational-Districts-Programme-An-Appraisal.pdf.','HindustanTimesLabs. India district boundaries GeoJSON. data/raw/india_districts.geojson; https://github.com/HindustanTimesLabs/shapefiles/tree/master/india.']),
    ('Annexes',['Crosswalk method: normalized exact names form provisional identifiers. Candidate fuzzy matches are suggestions only. Every approved name or parent/child mapping must be written to data/processed/district_name_overrides.csv and documented in data/processed/district_crosswalk.csv.','Indicator definitions available in NFHS-4 include “Women who are literate (%)”, “Women with 10 or more years of schooling (%)”, “Women age 20-24 years married before age 18 years (%)”, “Households using clean fuel for cooking”, and “Households using improved sanitation facility”. (Source: data/raw/nfhs_district-wise.csv.)','The logframe, theory of change, bilingual field tool, and sampling plan are included as separate editable outputs. Baselines and targets that depend on selected districts are labeled pending.'])]
    for idx,(heading,paragraphs) in enumerate(pages):
        d.add_heading(heading,0 if idx==0 else 1)
        for para in paragraphs:d.add_paragraph(para)
        if heading.startswith('RQ4'):
            d.add_picture(str(ROOT/'outputs'/'maps'/'women_literate_nfhs4.png'),width=Inches(5.6)); cap=d.add_paragraph('Figure 1. NFHS-4 women literacy baseline by district. Source and n are stated on the map.');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if heading.startswith('Annexes'):
            d.add_paragraph('Table 1. Core source indicator definitions currently identified in NFHS-4. Source: HindustanTimesLabs nfhs_district-wise.csv.')
            table(d,['Indicator','NFHS-4 source name','Round status'],[['Literacy','Women who are literate (%)','NFHS-5 common definition pending'],['Schooling','Women with 10 or more years of schooling (%)','NFHS-5 common definition pending'],['Early marriage','Women age 20-24 years married before age 18 years (%)','NFHS-5 common definition pending'],['Cooking fuel','Households using clean fuel for cooking','NFHS-5 common definition pending'],['Sanitation','Households using improved sanitation facility','NFHS-5 common definition pending']])
        if idx<len(pages)-1:d.add_page_break()
    path=ROOT/'outputs'/'reports'/'ZilaLens_Research_Report.docx';d.save(path);return path

def concept_note():
    d=Document();style_doc(d,'Concept Note');add_title(d,'Concept note draft | district selection pending');d.add_paragraph('Proposed 12-month CSR or donor-funded pilot in 3–5 low-performing Uttar Pradesh districts, to be selected after NFHS-5 analysis and field validation.');d.add_page_break()
    d.add_heading('1. Problem and rationale',1);d.add_paragraph('Public financial inclusion programmes have expanded formal access, while a woman’s account ownership does not by itself show whether she can use the account regularly, privately, and with control over transactions. ZilaLens proposes a district-focused pilot that will pair quantitative district evidence with field validation before final geographies are chosen. The project name combines “Zila”, Hindi for district, with a district lens on women’s financial inclusion.')
    d.add_paragraph('The NFHS-4 district compilation contains 59,241 district-indicator rows and 93 indicators according to the source repository documentation. (Source: data/raw/nfhs_district-wise.csv and data/raw/README.md.) NFHS-5 is required to identify the 3–5 lowest-performing UP districts and has not yet been downloaded. No UP ranking or district-specific problem estimate is asserted in this draft.')
    d.add_heading('2. Goal and objectives',1);bullet(d,'Goal: support women to access and use suitable financial services as part of stronger livelihood resilience.')
    bullet(d,'Objective 1: identify local account ownership, use, control, phone access, and service barriers through district data and field validation.')
    bullet(d,'Objective 2: improve practical use, digital safety, and access to responsive service points in selected districts.')
    bullet(d,'Objective 3: build a simple monitoring system that tracks women’s independent use and service quality, not enrolment alone.')
    d.add_heading('3. Intervention and results framework',1);table(d,['Workstream','Activities','Proposed measures'],[['Local service access','Map banks/BC points, hours, network, grievance route','Functional service points; issue resolution'],['Women’s financial capability','Small-group practice on balances, deposits, withdrawals, transfers and safety','Independent task completion; knowledge check'],['Peer and referral support','Train local resource persons; refer service failures and fraud concerns','Referral completion; participant feedback'],['Monitoring and learning','Baseline, quarterly checks, endline and adaptation','Account use, control and satisfaction measures']])
    d.add_page_break();d.add_heading('4. Implementation approach',1);d.add_paragraph('A local implementation partner would work with district administration, banks, BC networks, women’s collectives, and community organizations. The first two months would confirm district selection, partner capacity, safeguarding, and a service-point map. Months 3–9 would deliver community sessions and assisted practice. Months 10–12 would consolidate referrals, repeat measures, and document adaptations. No direct account opening targets should replace informed choice or service quality.')
    d.add_heading('5. Indicative budget',1);d.add_paragraph('Planning assumption: 4 districts for 12 months; values below are illustrative estimates in Indian rupees and require local costing and donor approval.')
    table(d,['Cost line','Calculation assumption','Indicative amount (₹ lakh)'],[['District facilitators','8 facilitators × 12 months × ₹45,000/month','43.2'],['Training and materials','40 sessions × ₹35,000','14.0'],['Travel and field supervision','4 districts × ₹2.5 lakh/month × 12 months','120.0'],['Baseline and endline','Survey design, fieldwork, analysis','35.0'],['Digital safety and referral support','Materials, helpdesk linkage, escalation','12.0'],['Programme management and audit','Central coordination, safeguarding, audit','24.0'],['Subtotal','','248.2'],['Contingency','5% of subtotal; planning assumption','12.4'],['Total','','260.6']])
    d.add_heading('6. Risks and mitigation',1);table(d,['Risk','Mitigation'],[['NFHS-5 file or boundary crosswalk delays district selection','Do not announce geographies before verified analysis; use a transparent selection note.'],['Women’s account-use interviews lack privacy','Use trained women enumerators, consent, private interviews and skip options.'],['BC/service point gaps or network outages','Map service availability; coordinate escalation with bank and telecom partners.'],['Digital fraud or harm','Teach safety, avoid requesting credentials, and provide verified reporting channels.'],['Programme measures reward account opening only','Track active use, control, problem resolution and participant experience.']])
    d.add_page_break();d.add_heading('7. Monitoring, evaluation and learning',1);d.add_paragraph('The attached logframe identifies outcome and service measures. Numeric baselines and targets will be set after district selection, sample design and power considerations. The proposed field tool validates independent account use and the gap between access and use. Pre/post comparisons would describe programme-associated change only; stronger causal claims would require a separately designed evaluation.')
    d.add_heading('8. SDG alignment',1);bullet(d,'SDG 1: strengthen household capacity to manage financial shocks.');bullet(d,'SDG 5: support women’s agency and control over financial decisions.');bullet(d,'SDG 8: improve access to useful financial services that can support livelihoods.');bullet(d,'SDG 10: prioritize underserved districts and monitor unequal access.')
    d.add_paragraph('Prepared by Kanishk Gupta | PGDM (Research & Business Analytics), Lal Bahadur Shastri Institute of Management, Delhi | '+GITHUB)
    p=ROOT/'outputs'/'reports'/'ZilaLens_Concept_Note.docx';d.save(p);return p

def policy_brief():
    d=Document();style_doc(d,'Policy Brief');sec=d.sections[0];sec.top_margin=Inches(.45);sec.bottom_margin=Inches(.45);add_title(d,'Interim evidence brief | comparative findings pending')
    d.add_heading('The district comparison is ready to run; NFHS-5 is the missing input',1);d.add_paragraph('The current evidence supports building the comparison and validating district definitions. It does not yet support claims about which districts improved or whether Aspirational Districts moved faster.')
    d.add_heading('What is verified so far',2)
    table(d,['Data readiness measure','Value','Source'],[['NFHS-4 district-indicator rows','59,241','HindustanTimesLabs nfhs_district-wise.csv'],['NFHS-4 district indicators','93','HindustanTimesLabs repository README'],['Exact-name ADP list matches in NFHS-4','96 of 112','NITI 112 list + NFHS-4 source; project crosswalk']])
    d.add_paragraph('The NITI 112-district PDF is published in 2022. Sixteen listed names are not found as exact NFHS-4 state–district names and need geography review. No outcome comparison has been computed.')
    d.add_page_break();d.add_heading('Three actions for a responsible district programme',1)
    bullet(d,'Place the official data.gov.in workbook at data/raw/NFHS_5_India_Districts_Factsheet_Data.xls and rerun the documented pipeline.')
    bullet(d,'Approve a parent/child district crosswalk before publishing rankings or selecting low-performing districts.')
    bullet(d,'Measure independent use, transaction control, phone access and service reliability alongside account ownership.')
    d.add_picture(str(ROOT/'outputs'/'maps'/'women_literate_nfhs4.png'),width=Inches(5.4));p=d.add_paragraph('Interim map: NFHS-4 women literacy baseline only; 563 district geometries matched. It does not show change. Source and boundary note appear on the map.');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph('Kanishk Gupta | '+GITHUB);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    path=ROOT/'outputs'/'reports'/'ZilaLens_Policy_Brief.docx';d.save(path);return path

def deck():
    prs=Presentation();prs.slide_width=PInches(13.333);prs.slide_height=PInches(7.5)
    prs.core_properties.author=AUTHOR;prs.core_properties.title=TITLE+' | '+SUB;prs.core_properties.subject='Interim evidence deck; NFHS-5 required for comparison'
    slides=[
      ('ZilaLens',SUB+'\nKanishk Gupta | '+GITHUB,'Interim edition, 18 October 2026.','title'),
      ('NFHS-5 is the missing input for all two-round findings','The official data.gov.in workbook resource returned HTTP 403 in this environment.','Catalog resource: “India Districts Factsheets of NFHS-5, 2019-2021”; exact destination is recorded in README and DATA_SOURCES.','text'),
      ('NFHS-4 provides a usable district baseline for selected indicators','59,241 district-indicator rows and 93 district indicators are present in the specified NFHS-4 compilation.','Source: HindustanTimesLabs nfhs_district-wise.csv and README. Counts describe the input file, not a population finding.','map1'),
      ('Improved sanitation varies across mapped districts at baseline','The interim map joins 563 district estimates to the available public district GeoJSON.','Source: NFHS-4 compilation and HindustanTimesLabs district GeoJSON. Boundary vintage is not confirmed as survey-equivalent.','map2'),
      ('Common indicators must be confirmed from both factsheet rounds','NFHS-4 includes literacy, schooling, early marriage, clean fuel and improved sanitation indicators.','Bank-account and mobile phone measures were not identified under the requested wording in the NFHS-4 compilation; exact NFHS-5 labels remain unverified.','text'),
      ('Name similarity cannot resolve parent and child districts by itself','Current crosswalk uses normalized names and retains fuzzy suggestions for review.','Document the selected parent/child mapping and any assumptions before calculating change.','text'),
      ('The 112-district roster needs a reviewed NFHS crosswalk','The NITI 2022 roster has 112 entries; 96 match NFHS-4 state–district names exactly.','Sixteen names remain unresolved by exact matching; source list and project crosswalk are cited in the report. This is a roster coverage statistic, not an outcome estimate.','text'),
      ('UP change rankings wait for the paired NFHS-5 estimates','Ghaziabad and Gautam Buddha Nagar can be compared with state and national benchmarks only after harmonization.','The interim UP map shows NFHS-4 baseline literacy only.','mapup'),
      ('Policy documents point to reach, connectivity and delivery capacity','The PMJDY brochure describes branch and Bank Mitra reach and notes telecom constraints in some locations.','PMJDY Mission Document/Brochure, physical pp. 18, 20 and 30; NITI/UNDP ADP appraisal, physical pp. 15, 20 and 28. Paraphrases are in the coding matrix.','text'),
      ('The intervention should measure use and control alongside access','Proposed pathway: reachable service → women’s independent use → financial options → resilience.','The theory of change states assumptions. No programme effect is claimed.','toc'),
      ('Field validation can test whether ownership translates into use','The bilingual XLSForm asks about recent transactions, decision control, phone access, service failures, help needs and barriers.','Consent gate and skip logic are included; pyxform conversion succeeded without warnings.','text'),
      ('District selection and final recommendations follow verified comparison','Provide the official workbook, then review the crosswalk and common indicators before selecting 3–5 UP districts.','Results remain descriptive. Aspirational District comparisons are not causal.','text'),
    ]
    for i,(title,body,note,visual) in enumerate(slides):
        s=prs.slides.add_slide(prs.slide_layouts[6]);bg=s.background.fill;bg.solid();bg.fore_color.rgb=PRGB(248,250,252)
        if i==0:
            box=s.shapes.add_textbox(PInches(.8),PInches(1.4),PInches(11.7),PInches(1.0));tf=box.text_frame;tf.text=title;tf.paragraphs[0].font.size=PPt(36);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=PRGB(23,54,93)
            box=s.shapes.add_textbox(PInches(.8),PInches(2.55),PInches(10.8),PInches(1.6));tf=box.text_frame;tf.text=body;tf.paragraphs[0].font.size=PPt(19);tf.paragraphs[0].font.color.rgb=PRGB(22,140,140)
            add_footer(s,i+1);s.notes_slide.notes_text_frame.text=note;continue
        titlebox=s.shapes.add_textbox(PInches(.65),PInches(.4),PInches(12),PInches(1.0));p=titlebox.text_frame.paragraphs[0];p.text=title;p.font.size=PPt(26);p.font.bold=True;p.font.color.rgb=PRGB(23,54,93)
        bodybox=s.shapes.add_textbox(PInches(.8),PInches(1.65),PInches(11.7),PInches(2.0));p=bodybox.text_frame.paragraphs[0];p.text=body;p.font.size=PPt(23);p.font.color.rgb=PRGB(38,50,56)
        if visual.startswith('map'):
            img='women_literate_nfhs4.png' if visual=='map1' else ('women_literate_up_nfhs4.png' if visual=='mapup' else 'improved_sanitation_nfhs4.png')
            s.shapes.add_picture(str(ROOT/'outputs'/'maps'/img),PInches(8.0),PInches(3.05),height=PInches(3.8))
        elif visual=='toc':s.shapes.add_picture(str(ROOT/'outputs'/'charts'/'theory_of_change.png'),PInches(.9),PInches(3.3),width=PInches(11.6))
        noteBox=s.shapes.add_textbox(PInches(.8),PInches(5.9),PInches(6.6 if visual.startswith('map') else 11.7),PInches(.9));p=noteBox.text_frame.paragraphs[0];p.text=note;p.font.size=PPt(10);p.font.color.rgb=PRGB(91,101,115)
        add_footer(s,i+1);s.notes_slide.notes_text_frame.text=note
    path=ROOT/'outputs'/'deck'/'ZilaLens_Deck.pptx';prs.save(path);return path

def add_footer(slide,num):
    box=slide.shapes.add_textbox(PInches(.65),PInches(7.12),PInches(12),PInches(.2));p=box.text_frame.paragraphs[0];p.text=f'ZilaLens  |  Kanishk Gupta                                               {num:02d}';p.font.size=PPt(9);p.font.color.rgb=PRGB(91,101,115)

def build_all():
    (ROOT/'outputs'/'reports').mkdir(parents=True,exist_ok=True)
    report();concept_note();policy_brief();deck()

if __name__=='__main__':build_all()
